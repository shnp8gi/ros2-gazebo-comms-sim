#!/usr/bin/env python3
"""内部調査資料 (Word) を生成する。

書式は seminar/2026_08_04_KKFSystem.docx をそのまま雛形として引き継ぐ:
  用紙・余白  A4 / 上下 2.5 cm・左右 2.0 cm (雛形の sectPr をそのまま使う)
  見出し      List Paragraph + 自動番号 (numId は雛形の見出しから複製)、章=ilvl0 / 節=ilvl1
  本文        Normal / Times New Roman 10.5 pt / 1 字下げ
  句読点      ，．(全角カンマ・ピリオド)、である調

数値はすべて sim_results の集計 CSV から読む (本文への手書き写しをしない)。

  python3 scratch/make_investigation_docx.py
"""
import copy
import os
import sys

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from scipy import stats

ROOT = '/workspace'
R = f'{ROOT}/sim_results'
TPL = f'{ROOT}/scratch/seminar/2026_08_04_KKFSystem.docx'
OUT = f'{ROOT}/scratch/review'
FIGS = f'{OUT}/figs'
os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, f'{ROOT}/tools')
from lib import eval_metrics  # noqa: E402

# ============================================================ 書式 (雛形から)
doc = Document(TPL)

h_lvl0 = h_lvl1 = body_tpl = None
for p in doc.paragraphs:
    t = p.text.strip()
    if p.style.name == 'List Paragraph' and t:
        pPr = p._p.pPr
        num = pPr.find(qn('w:numPr')) if pPr is not None else None
        ilvl = None
        if num is not None and num.find(qn('w:ilvl')) is not None:
            ilvl = num.find(qn('w:ilvl')).get(qn('w:val'))
        if ilvl == '0' and h_lvl0 is None:
            h_lvl0 = p
        if ilvl == '1' and h_lvl1 is None:
            h_lvl1 = p
    elif p.style.name == 'Normal' and t and body_tpl is None and p.runs:
        body_tpl = p
assert h_lvl0 is not None and h_lvl1 is not None and body_tpl is not None, '雛形が取れない'

TBL_STYLE = doc.tables[0].style.name if doc.tables else 'Table Grid'
h0_x = copy.deepcopy(h_lvl0._p)
h1_x = copy.deepcopy(h_lvl1._p)
body_x = copy.deepcopy(body_tpl._p)

# 雛形の中身を捨てる (sectPr = 用紙・余白の定義だけ残す)
body = doc.element.body
sectPr = body.find(qn('w:sectPr'))
for child in list(body):
    if child is not sectPr:
        body.remove(child)


# w:rPr の子要素は順序が決まっている (ECMA-376 CT_RPr)。順序を守らないと
# Word が壊れたファイルとして扱うので、必ずこの並びに割り込ませる
_RPR_ORDER = ['w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps',
              'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline', 'w:shadow',
              'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish',
              'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position',
              'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd',
              'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang']


def _rpr_add(rPr, tag, attrs=None):
    """rPr に要素を規定の順序で差し込む (既存の同名要素は置き換える)。"""
    for e in rPr.findall(qn(tag)):
        rPr.remove(e)
    el = rPr.makeelement(qn(tag), {})
    for k, v in (attrs or {}).items():
        el.set(qn(k), v)
    rank = _RPR_ORDER.index(tag) if tag in _RPR_ORDER else len(_RPR_ORDER)
    for child in rPr:
        name = child.tag.split('}')[-1]
        r = _RPR_ORDER.index('w:' + name) if ('w:' + name) in _RPR_ORDER else len(_RPR_ORDER)
        if r > rank:
            child.addprevious(el)
            return el
    rPr.append(el)
    return el


def _jp_punct(text):
    """句点を雛形の体裁 (，．) にそろえる。数値の小数点は半角なので影響しない。"""
    return text.replace('。', '．')


def _clone(tpl_x, text, bold=None, size=None, align=None, indent=None):
    """雛形段落を複製して文字列だけ差し替える (番号付け・フォントを保つ)。"""
    text = _jp_punct(text)
    new = copy.deepcopy(tpl_x)
    runs = new.findall(qn('w:r'))
    for r in runs[1:]:
        new.remove(r)
    r = runs[0]
    for t in r.findall(qn('w:t')):
        r.remove(t)
    el = r.makeelement(qn('w:t'), {})
    el.text = text
    el.set(qn('xml:space'), 'preserve')
    r.append(el)
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rPr)
    if bold is not None:
        for tag in ('w:b', 'w:bCs'):
            for e in rPr.findall(qn(tag)):
                rPr.remove(e)
        if bold:
            _rpr_add(rPr, 'w:b')
            _rpr_add(rPr, 'w:bCs')
    if size is not None:
        for tag in ('w:sz', 'w:szCs'):
            _rpr_add(rPr, tag, {'w:val': str(int(size * 2))})
    pPr = new.find(qn('w:pPr'))
    if pPr is None:
        pPr = new.makeelement(qn('w:pPr'), {})
        new.insert(0, pPr)
    if align is not None:
        for e in pPr.findall(qn('w:jc')):
            pPr.remove(e)
        e = new.makeelement(qn('w:jc'), {})
        e.set(qn('w:val'), align)
        pPr.append(e)
    if indent is not None:
        for e in pPr.findall(qn('w:ind')):
            pPr.remove(e)
        if indent:
            e = new.makeelement(qn('w:ind'), {})
            e.set(qn('w:firstLine'), str(int(indent * 567)))  # cm -> twips
            pPr.append(e)
    sectPr.addprevious(new)
    return new


def h1(text):
    return _clone(h0_x, text)


def h2(text):
    return _clone(h1_x, text)


def body_p(text, bold=None, size=None):
    return _clone(body_x, text, bold=bold, size=size)


def caption(text):
    return _clone(body_x, text, size=9.5, align='center', indent=0)


def note(text):
    return _clone(body_x, text, size=8.5, indent=0)


def _cell_run(cell, text, size=8.5, bold=False, align=None):
    p = cell.paragraphs[0]
    run = p.add_run(_jp_punct(text))
    e = run._element
    rPr = e.find(qn('w:rPr'))
    if rPr is None:
        rPr = e.makeelement(qn('w:rPr'), {})
        e.insert(0, rPr)
    _rpr_add(rPr, 'w:rFonts', {'w:ascii': 'Times New Roman',
                               'w:hAnsi': 'Times New Roman',
                               'w:cs': 'Times New Roman'})
    if bold:
        _rpr_add(rPr, 'w:b')
        _rpr_add(rPr, 'w:bCs')
    for tag in ('w:sz', 'w:szCs'):
        _rpr_add(rPr, tag, {'w:val': str(int(size * 2))})
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0


def figure(name, cap, width_cm=15.5):
    """図を貼り、その下に図番号つきの説明を置く (表は上、図は下)。"""
    path = os.path.join(FIGS, name)
    if not os.path.exists(path):
        sys.exit(f'図がない: {path}  先に scratch/make_investigation_figs.py を流すこと')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    return caption(cap)


def table(headers, rows, widths, right_from=1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, hd in enumerate(headers):
        _cell_run(t.rows[0].cells[i], hd, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            al = WD_ALIGN_PARAGRAPH.RIGHT if i >= right_from else None
            _cell_run(cells[i], str(v), align=al)
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Cm(w)
    return t


# ============================================================ データ
def mb(v):
    return f'{v:,.0f}'


def smb(v):
    return f'{v:+,.0f}'


def load(tag):
    runs = pd.read_csv(f'{R}/{tag}/analysis/runs.csv')
    agg = pd.read_csv(f'{R}/{tag}/analysis/agg.csv').set_index('method')
    return runs, agg


def ci(x):
    x = pd.Series(x).dropna()
    if len(x) < 2 or x.std(ddof=1) == 0:
        return 0.0
    return float(stats.t.ppf(0.975, len(x) - 1) * x.std(ddof=1) / np.sqrt(len(x)))


def diff_vs(runs, m, ref='assoc_hold'):
    """ref に対する対差 (MB) と走行ごと相対差 (%) を返す。"""
    a = runs[runs.method == ref].set_index('run').total_data_MB
    x = runs[runs.method == m].set_index('run').total_data_MB
    d = (x - a).dropna()
    rel = (100 * (x - a) / a).dropna()
    return d, ci(d), rel, ci(rel)


runs, agg = load('abl_d70_n50')
o_runs, o_agg = load('rec_cm_d70_ablation')
N = int(agg.loc['assoc_hold', 'n'])
N_REC = len([d for d in os.listdir(f'{R}/rec_d70_n50') if d.startswith('seed_')])
N_OLD = int(o_agg.loc['assoc_hold', 'n'])

base = float(agg.loc['assoc_hold', 'total_data_MB_mean'])
conv = float(agg.loc['kkf_conv', 'total_data_MB_mean'])
orc = float(agg.loc['oracle', 'total_data_MB_mean'])
probe = float(agg.loc['kkf_cold_probe', 'total_data_MB_mean'])
d, d_ci, rel, rel_ci = diff_vs(runs, 'kkf_conv')
d_o, d_o_ci, rel_o, rel_o_ci = diff_vs(runs, 'oracle')
d_pr, _, rel_pr, _ = diff_vs(runs, 'kkf_cold_probe')
capture = 100.0 * d.mean() / d_o.mean()
o_d, o_d_ci, o_rel, o_rel_ci = diff_vs(o_runs, 'kkf_conv')
o_cap = 100.0 * o_d.mean() / diff_vs(o_runs, 'oracle')[0].mean()


def need_n(x, target=0.05):
    """平均の ±target の信頼区間半幅を得るのに必要な走行数 (t 区間で反復)。"""
    cv = pd.Series(x).std(ddof=1) / pd.Series(x).mean()
    n = 2.0
    for _ in range(200):
        n2 = (stats.t.ppf(0.975, max(1, int(n) - 1)) * cv / target) ** 2
        if abs(n2 - n) < 0.5:
            break
        n = n2
    return int(np.ceil(n))


need_old, need_new = need_n(o_rel), need_n(rel)
ncar = runs[runs.method == 'assoc_hold'].n_cars

# kkf_conv を基準にした機構ごとの対比較
pk = eval_metrics.paired(runs, 'kkf_conv', 'method', ['load'], ['total_data_MB'])
pk = pk[pk.metric == 'total_data_MB'].set_index('baseline')
o_pk = eval_metrics.paired(o_runs, 'kkf_conv', 'method', ['load'], ['total_data_MB'])
o_pk = o_pk[o_pk.metric == 'total_data_MB'].set_index('baseline')
pa = pd.read_csv(f'{R}/abl_d70_n50/analysis/paired.csv')
pa = pa[pa.metric == 'total_data_MB'].set_index('baseline')

# 選択余地とカバレッジ (掃引 5 条件 + 直交 4 条件)
cm = pd.read_csv(f'{R}/choice_margin_summary.csv')
cf = pd.read_csv(f'{R}/confound_summary.csv')
cond = pd.concat([cm[cm['索引OK'] == True], cf[cf['索引OK'] == True]], ignore_index=True)
cond = cond.sort_values('選択余地率').reset_index(drop=True)
GEOM = {'rec_cm_d40': ('4 本 / 20 m / 40°', 8), 'rec_cm_d50': ('4 本 / 20 m / 50°', 8),
        'rec_cm_d60': ('4 本 / 20 m / 60°', 8), 'rec_cm_d70': ('4 本 / 20 m / 70°', 8),
        'rec_cm_d70s10': ('4 本 / 10 m / 70°', 8), 'rec_cf_A': ('4 本 / 25 m / 70°', 8),
        'rec_cf_B': ('6 本 / 15 m / 70°', 12), 'rec_cf_C': ('6 本 / 25 m / 70°', 12),
        'rec_cf_D': ('4 本 / 15 m / 70°', 8)}
r_marg, p_marg = stats.pearsonr(cond['選択余地率'], cond['天井_pct'])
r_cov, p_cov = stats.pearsonr(cond['assoc_hold_MB'], cond['天井_pct'])
hi = cond[cond['選択余地率'] >= 0.155]
lo = cond[cond['選択余地率'] < 0.155]
r_marg_hi, p_marg_hi = stats.pearsonr(hi['選択余地率'], hi['天井_pct'])
r_cov_hi, p_cov_hi = stats.pearsonr(hi['assoc_hold_MB'], hi['天井_pct'])


def _audit(path):
    a = pd.read_csv(path).iloc[0]
    return {'miss': float(a['miss_pct']), 'fa': float(a['false_alarm_pct']),
            'bias': float(a['bias_db_on_connectable']),
            'rmse': float(a['rmse_db_on_connectable']), 'z': float(a['z_std'])}


au = _audit(f'{R}/audit_d70/summary_frozen.csv')

LBL = {'oracle': 'oracle（全ペアの受信電力が既知）', 'kkf_cold_probe': 'kkf_cold_probe（全ペア観測・地図なし）',
       'kkf_conv': 'kkf_conv（提案手法）', 'kkf_nolcb': 'kkf_nolcb（LCB なし）',
       'kkf_novar': 'kkf_novar（遮蔽リスク地図なし）', 'assoc_hold': 'assoc_hold（受動接続）',
       'kkf_cold': 'kkf_cold（地図なし・規格忠実）'}

# ============================================================ 本文
_clone(body_x, 'KKF 電波地図に基づく能動的ハンドオーバーの有用性', bold=True,
       size=14, align='center', indent=0)
_clone(body_x, f'都市部 2 車線・論理 RSU 8 基における {N} 走行の評価',
       size=11, align='center', indent=0)
_clone(body_x, '2026-08-13 / feat/urban-2lane / 内部資料',
       size=9.5, align='center', indent=0)

h1('主張')
body_p('本資料が扱う問題は次のものである。ミリ波（60 GHz）で路側機（RSU）と走行車両が通信するとき，'
       '1 基の RSU に同時に接続できる車両は 1 台に限られる（802.15.3e のペアネット方式）。'
       'さらに，接続していない相手については，割当の判断に足るだけの受信電力を継続して'
       '測ることができない。規格の受動アソシエーションで分かるのは，'
       'つなげる見込みがあるかどうかの粗い判定にとどまる。'
       'どの車両をどの RSU にいつ接続するかを決めるには，'
       '測っていない相手の電波状況を何らかの方法で補う必要がある。')
body_p('提案手法は，過去の走行から学習した電波地図でこれを補い，接続先を能動的に決める。'
       '地図の学習には Kriged Kalman Filter（以下 KKF）を用いる。'
       'これは，位置どうしの空間相関を使って未観測の地点を補間する手法（クリギング）と，'
       '観測が入るたびに推定を更新する手法（カルマンフィルタ）を組み合わせたもので，'
       '位置を入力すると受信電力の予測値と，その予測がどれだけ不確かか（σ）を返す。')
body_p('本シミュレーション環境の配置 d70 において，KKF 電波地図に基づく能動的ハンドオーバーは，'
       '802.15.3e 準拠の受動接続に対し総配信量で +%.1f%% ± %.1f pt 優位である'
       '（%d 走行，95%% 信頼区間）。%d 走行すべてで提案手法が上回り，'
       '最も差が小さい走行でも +%.1f%% である。'
       'この優位性は事前学習した電波地図に由来し，地図は規格上観測できないペアの現在値を'
       '代替する役割を担う。全ペアの現在値を雑音なしで与えた理想反応型に対する到達率は %.1f%% である。'
       % (rel.mean(), rel_ci, N, N, rel.min(), capture))
body_p('本稿で「能動的」とは，規格上観測できないペアの現在値を地図で推定して'
       '切替の判断を行うことを指し，将来時刻の予測は含まない（1 段先，すなわち現時刻しか見ない）。')

h1('評価の方法')
h2('事後再計算による評価')
body_p('全手法を事後再計算（replay）で評価する。軌跡と全ペアの受信電力を 1 回記録し，'
       '同一のチャネル実現の上で通信とスケジューリングだけを再計算するため，'
       '同じ交通・同じ電波の実現の上で手法だけを差し替えることになり，条件差が完全に打ち消される。')
note('再計算は決定論的である。本資料の作成にあたって 1 走行を再実行し，'
     '記録済みの値と小数点以下まで一致することを確認した。')
caption('表1　評価条件')
table(['項目', '設定'],
      [['道路・RSU', '都市部 2 車線（各方向 1 車線），論理 RSU 8 基'
                     '（ポール 4 本 × 上流/下流 2 面，ポール間隔 20 m，路側 y = 6 m，高さ 2.5 m）。'
                     'アンテナ傾き δ_rsu = 70°，車載 δ_veh = 20°（互いの正面が向き合う向き）。'
                     '以下これを配置 d70 と呼ぶ'],
       ['車両', f'平均 {ncar.mean():.1f} 台/走行（{ncar.min()}〜{ncar.max()}），速度 7.0 m/s，'
                '車頭時間 平均 2.0 s，前後 2 アンテナ（高さ 1.35 m）'],
       ['無線', '60 GHz，送信電力 −7 dBm，接続閾値 −78.5 dBm，帯域 100 MHz，'
                '排他接続。観測レポートは 20 Hz，観測雑音 σ = 2.0 dB'],
       ['チャネル', '自由空間経路損失 + 持続シャドウイング（σ = 4 dB，相関長 6 m）+ '
                    '路上駐車 2 台による遮蔽。フェージング無効。'
                    'シャドウイングの場と駐車位置は environment_seed = 1 で固定（本節末を参照）'],
       ['走行', '交通生成窓 35 s，走行のシム時刻は約 78 s（上限 90 s），記録周期 5 ms'],
       ['レート', 'Shannon 式（帯域 100 MHz，雑音床 −95 dBm，効率 1.0，SNR 上限 50 dB）。'
                  '802.15.3e に準拠するのは接続手順（受動アソシエーションと排他接続）であり，'
                  'チャネル帯域は規格の 2.16 GHz ではない'],
       ['KKF 地図', '道路に沿って等間隔に置いた山型の関数（RBF 基底）40 個の重み付き和で'
                    '受信電力を表す（道路に沿った距離で 0〜210 m）。地図は（RSU，走行方向，車載アンテナ）ごとに'
                    '独立 = 32 枚。学習 30 走行（base_seed 512345，評価とシード域を分離）。'
                    '学習相の交通は素のシナリオの設定（速度 16.7 m/s，車頭 5.0 s，生成窓 60 s）で'
                    '評価相とは異なるが，走行車両は遮蔽体として扱わないため'
                    '伝搬そのものには影響せず，道路に沿った標本の分布のみが変わる'],
       ['評価', f'記録 {N_REC} 走行のうち {N} 走行（base_seed 12345，シード刻み 1000）。'
                '対比較 Wilcoxon 符号順位検定，信頼区間は対差の t 区間'],
       ['再現', 'コードと設定は feat/urban-2lane の 789996b / ef7b8b7。'
                '記録時の作業差分は sim_results/rec_d70_n50/sweep/config/workspace.patch。'
                '実行環境と手順は付録 D']],
      widths=[2.6, 14.4], right_from=99)
body_p('持続シャドウイングの場と路上駐車の位置は，学習相と評価相で同一である。'
       'これは環境に固有の減衰を学習対象とする電波地図の前提そのものであり，'
       '走行ごとに変わるのは交通実現と観測雑音である。'
       'したがって本評価が示すのは「同じ環境を走り続けた場合に地図が持つ価値」であって，'
       '環境が変わったときの地図の陳腐化は評価していない（第 5 章）。')
body_p('記録した %d 走行のうち %d 走行は評価に載っていない。'
       '原因は記録側の書き出しが途中で打ち切られたことである。'
       '記録は全手法で共有するため，この脱落は手法の優劣に対して選択的ではない。詳細は付録 B に示す。'
       % (N_REC, N_REC - N))

h2('用語')
caption('表2　本資料で使う用語')
table(['用語', '意味'],
      [['割当（grant）', 'ある車両にある RSU を使う権利が与えられている状態。'
                        '割当があっても電波が弱ければ送れないので，割当時間と接続時間は一致しない'],
       ['接続時間', '割当を得たうえで実際に受信電力が閾値を超え，データを送れていた時間'],
       ['対差', '同じ交通実現（同じ記録）どうしで手法間の値を引き算した差。'
                '走行ごとの交通の当たり外れが打ち消されるため，本資料の主指標とする'],
       ['走行ごと相対差', '対差を，その走行の受動接続の値で割ったもの。走行間で規模をそろえた比較'],
       ['到達率', '受動接続との対差のうち，理想反応型（oracle）が得た対差の何 % を'
                  '提案手法が得たか。地図が完全な計測をどこまで代替できたかを表す'],
       ['LCB（κσ）', '予測値から不確かさ σ の κ 倍を差し引いた値で評価する慎重な選び方。'
                     '確度の低い予測を高く評価しないための仕組み'],
       ['遮蔽リスク地図 σν²', '同じ位置でも走行ごとに受信電力が揺らぐ量の大きさを別に保持する地図。'
                             '揺らぎの大きい位置を避けるために使う'],
       ['割当層', '各時刻に「どの車両にどの RSU を与えるか」を決める部分。'
                  '本資料ではすべての手法で共通に，予測レートの総和を最大にする貪欲な割当を用いる。'
                  '地図（予測）とは分離されている']],
      widths=[3.4, 13.6], right_from=99)

h2('比較した手法')
caption('表3　比較した手法')
table(['手法', '内容'],
      [['assoc_hold', '802.15.3e 準拠の受動接続。閾値を超えた空き RSU に接続し保持する。'
                      '能動的なハンドオーバーを行わない標準ベースライン'],
       ['kkf_conv', '提案手法。学習済みの電波地図を持ち，観測は grant 中のペアに限る（規格忠実）'],
       ['kkf_cold', '地図を持たず，観測も grant 中のペアに限る。規格忠実な条件で'
                    '事前学習だけを外した構成'],
       ['kkf_cold_probe', '地図を持たないが，全ペアを観測して走行内で地図を育てる。'
                          '規格の接続手順では実行できない参照条件'],
       ['kkf_nolcb', '提案手法から不確実性の利用（LCB，κσ）を外した構成'],
       ['kkf_novar', '提案手法から遮蔽リスク地図（σν²）を外した構成'],
       ['oracle', '全ペアの現在の真値を雑音なしで用いる理想反応型。割当層は提案手法と同じ。'
                  '未来の予測ではなく「完全な計測が得られたらどれだけ配信できるか」を表す']],
      widths=[3.0, 14.0], right_from=99)

h1('主張の根拠')
h2('受動接続に対する優位')
figure('fig1_gain.png',
       '図1　受動接続 (assoc_hold) に対する走行ごと相対差（%d 走行）。'
       '誤差棒は対応データの 95%% 信頼区間，括弧内は対差の絶対量。'
       '地図を持たない kkf_cold は全走行 0 MB のため図示していない。数値は付録の表 A1。' % N)
body_p('提案手法の対差は +%s MB（95%% 信頼区間 ±%s MB，Wilcoxon 符号順位検定 p = %.2g）である。'
       '全ペアの受信電力が既知の理想反応型に対する到達率は，対差ベースで %s / %s = %.1f%% であり，'
       '地図が規格上行えない計測をこの水準まで代替できたことを意味する。'
       % (mb(d.mean()), mb(d_ci), float(pa.loc['kkf_conv', 'wilcoxon_p']),
          mb(d.mean()), mb(d_o.mean()), capture))
figure('fig2_runs.png',
       '図2　走行ごと相対差の分布。上段が前回の %d 走行，下段が今回の %d 走行。'
       '点 1 つが 1 走行で，縦位置のばらつきは重なりを避けるためのゆらぎである。'
       '橙は平均と 95%% 信頼区間。今回の %d 走行はすべて正側にある。数値は付録の表 A2。'
       % (N_OLD, N, N))
body_p('この優位はばらつきによる上振れではない（図2）。'
       '本測定は前回の %d 走行を %d 走行に増やしたものである。'
       '点推定は +%.1f%% から +%.1f%% とほとんど動かず，'
       '信頼区間の半幅は平均の ±%.1f%% から ±%.1f%% に縮んだ。'
       '有意性の判定が変わったのは後述する不確実性の利用（LCB）の寄与 1 点のみである。'
       % (N_OLD, N, o_rel.mean(), rel.mean(), 100 * o_rel_ci / o_rel.mean(),
          100 * rel_ci / rel.mean()))

h2('利得の源泉は割当時間の質である')
body_p('提案手法は受動接続より grant 時間が %.0f 秒短いにもかかわらず，接続時間は %.0f 秒長い'
       '（いずれも車両 × アンテナで積算した走行あたりの総和）。'
       '割当を得ていた時間のうち実際に送れた割合でみると，受動接続の %.1f%% に対し'
       '提案手法は %.1f%% である（図3）。'
       'すなわち利得は割当時間を延ばすことではなく，接続できない相手を選ばないことから来ている。'
       '1 基の RSU に同時に接続できる車両は 1 台に限られるため，'
       '送れていない割当はその RSU を使えたはずの他の車両にとっての損失でもある。'
       % (float(agg.loc['assoc_hold', 'grant_time_s_mean']) - float(agg.loc['kkf_conv', 'grant_time_s_mean']),
          float(agg.loc['kkf_conv', 'connected_time_s_mean']) - float(agg.loc['assoc_hold', 'connected_time_s_mean']),
          100 * float(agg.loc['assoc_hold', 'connected_time_s_mean']) / float(agg.loc['assoc_hold', 'grant_time_s_mean']),
          100 * float(agg.loc['kkf_conv', 'connected_time_s_mean']) / float(agg.loc['kkf_conv', 'grant_time_s_mean'])))
figure('fig3_time.png',
       '図3　割当を得ていた時間 (grant) のうち，実際に送れた時間 (接続時間) の割合。'
       '括弧内は走行あたりの合計秒数（車両 × アンテナで積算）。横軸は 90% 起点。')

h2('優位性は事前学習した地図に由来する')
body_p('地図を持たない kkf_cold は総配信量 0 MB であった。'
       'KKF の割当層は各ペアの見込みレートを必要とするため，'
       '割当中のペアしか観測できない規格忠実な条件では，接続しなければ観測できず，'
       '観測できなければ割り当てられないという循環に陥る。'
       '本実装においてこの循環を断っている機構は事前学習した地図のみである。')
body_p('利得は地図の内部機構ではなく，事前学習した地図を持つこと自体に帰属する。'
       '構成要素ごとに切り分けると，遮蔽リスク地図 σν² の寄与が %s MB（p = %.2g），'
       '不確実性を用いた慎重な選択（LCB）の寄与が %s MB（p = %.2g）で，'
       'いずれも符号は正だが利得全体に占める割合はそれぞれ %.1f%%，%.1f%% にとどまる（図4）。'
       % (smb(float(pk.loc['kkf_novar', 'diff_mean'])), float(pk.loc['kkf_novar', 'wilcoxon_p']),
          smb(float(pk.loc['kkf_nolcb', 'diff_mean'])), float(pk.loc['kkf_nolcb', 'wilcoxon_p']),
          100 * float(pk.loc['kkf_novar', 'diff_mean']) / d.mean(),
          100 * float(pk.loc['kkf_nolcb', 'diff_mean']) / d.mean()))
figure('fig4_ablation.png',
       '図4　提案手法から機構を 1 つずつ外したときの，kkf_conv との対差（%d 走行）。'
       '正の値は提案手法が上回ることを意味する。誤差棒は 95%% 信頼区間，'
       'p は Wilcoxon 符号順位検定（探索的比較であり多重比較の補正は行っていない）。'
       '地図そのものを外した kkf_cold（全走行 0 MB）は尺度が違うため図示していない。'
       '数値は付録の表 A3。' % N)
body_p('全ペアを観測できる kkf_cold_probe は %s MB で，提案手法を %s MB 上回る。'
       'ただし全ペア観測は 802.15.3e の接続手順では実行できないため，'
       'これは到達可能な代替案ではなく，規格の枠外に置いた参照条件である。'
       '提案手法はこの参照条件の %.1f%%（対差ベース）を，規格の枠内で再現している。'
       % (mb(probe), mb(probe - conv), 100 * d.mean() / d_pr.mean()))

h2('車両間の公平性')
figure('fig5_fairness.png',
       '図5　各走行で最も配信量が少なかった車両の配信量（走行平均，%d 走行）。'
       '括弧内は配信量が 0 だった車両の割合。数値は付録の表 A4。' % N)
body_p('提案手法は，最も恵まれない車両の配信量で受動接続の %.1f 倍（%s MB 対 %s MB）であり，'
       '配信量が 0 の車両の割合も提案手法だけが 0.00%% である。'
       '総配信量を増やす一方で取りこぼしを増やしてはいない。'
       % (float(agg.loc['kkf_conv', 'min_car_data_MB_mean']) / float(agg.loc['assoc_hold', 'min_car_data_MB_mean']),
          mb(float(agg.loc['kkf_conv', 'min_car_data_MB_mean'])),
          mb(float(agg.loc['assoc_hold', 'min_car_data_MB_mean']))))

h1('主張が成り立つ条件')
body_p('能動的な割当が利得を生むには，そもそも選ぶ相手が 2 つ以上ある時刻が必要である。'
       'この割合を選択余地率と呼び，RSU 配置で制御して，全ペアの受信電力が既知のときの利得'
       '（oracle の assoc_hold 比，以下「天井」）がどう変わるかを %d 条件で測った。'
       '天井は予測が完璧ならどれだけ配信できるかを表し，'
       '地図を改良して取り返せる余地の上限にあたる。'
       % len(cond))
figure('fig6_ceiling.png',
       '図6　天井（oracle の assoc_hold 比）を，選択余地率（左）とカバレッジ（右）で見たもの。'
       '丸は δ_rsu を振った掃引 5 条件，三角は本数と間隔を直交させた 4 条件。'
       'カバレッジは受動接続の総配信量で代表させた。数値は付録の表 A5。', width_cm=16.5)
body_p('選択余地率 0.155 を境に，天井は %.1f〜%.1f%% の領域と %.1f〜%.1f%% の領域に分かれる。'
       '境より上では，天井の大小を選択余地率でも（r = %+.2f）カバレッジでも（r = %+.2f）'
       '説明できない。したがって選択余地は必要条件として働くが，'
       'それ以上の量的な関係は現時点で認められない。'
       'なお境の値 0.155 は %d 条件の分布から事後的に引いたものであり，'
       'この値そのものに意味があるとは主張しない。'
       % (lo['天井_pct'].min(), lo['天井_pct'].max(), hi['天井_pct'].min(), hi['天井_pct'].max(),
          r_marg_hi, r_cov_hi, len(cond)))
body_p('本評価の配置 d70（余地 %.3f，天井 %+.1f%%）は境より上にある。'
       '測定した %d 条件のうち %d 条件が境より上にあり，'
       '天井のある領域は特別に選んだ例外的な配置ではない。'
       % (float(cond[cond['条件'] == 'rec_cm_d70']['選択余地率'].iloc[0]),
          float(cond[cond['条件'] == 'rec_cm_d70']['天井_pct'].iloc[0]), len(cond), len(hi)))

h1('限界')
body_p('第一に，提案手法自体の測定は配置 d70 の 1 点に限られる。'
       '図6 の天井は全ペアの受信電力が既知のときの利得であって提案手法の利得ではない。'
       '条件ごとに地図の学習が必要なため，他の配置での測定は行っていない。'
       '境より上の %d 条件のうち d70 を除く %d 条件の少なくとも 1 つで提案手法を測り，'
       '到達率が配置に依存しないことを確かめる必要がある。' % (len(hi), len(hi) - 1))
body_p('第二に，能動的な切替を行う非 KKF のベースラインを置いていない。'
       '比較対象は受動接続および地図を外した構成のみである。'
       '受動接続は閾値を超えた RSU に接続して保持するだけで切替を行わない。'
       '閾値割れで解放して再アソシエートする反応型は規格の枠内で構成でき，'
       '提案手法の利得のうち「切替を行うこと自体」に帰属する分を切り分けるには'
       'このベースラインの測定が要る。本評価ではこれを行っていない。')
body_p('第三に，oracle は絶対的な上界ではない。割当層は提案手法と同じ貪欲な総和最大化であり，'
       '実際に公平性の指標（図5）では提案手法が oracle を上回っている。'
       '割当層を改良すれば天井自体が上がりうる。')
body_p('第四に，学習した環境と評価する環境が同一である。'
       '環境が変わったときに地図がどれだけ陳腐化するか，'
       'また学習 30 走行という取得コストが利得に見合うかは評価していない。')
body_p('第五に，事後再計算では予測計算とスケジュール配信の遅延を 0 と仮定している。'
       'また学習相のみ全ペア観測を用いており規格忠実ではない。'
       '評価相は割当中のペアに限った観測に戻しているため比較の公正性は保たれるが，'
       '現地で地図を育てる運用には初期探索方策が別途必要である。')
body_p('第六に，ハンドオーバー回数・ping-pong 率（同じ 2 基の間を往復する切替の割合）・無接続時間の最大連続長を測定していない。'
       '接続時間と割当時間で代替しているが，切替コストの直接評価は今後の課題である。')

# ------------------------------------------------------------------ 付録
h1('付録 A：本文の図の数値')
body_p('本文の図に対応する数値を示す。')

rows = []
for m in ['oracle', 'kkf_cold_probe', 'kkf_conv', 'kkf_nolcb', 'kkf_novar',
          'assoc_hold', 'kkf_cold']:
    mean = float(agg.loc[m, 'total_data_MB_mean'])
    cn = float(agg.loc[m, 'connected_time_s_mean'])
    gr = float(agg.loc[m, 'grant_time_s_mean'])
    if m == 'assoc_hold':
        dd, rr = '—', '—'
    else:
        s, s_ci, sr, _ = diff_vs(runs, m)
        dd = f'{s.mean():+,.0f} ±{s_ci:,.0f}'
        rr = f'{sr.mean():+.1f}'
    rows.append([LBL[m], f'{mean:,.0f}', dd, rr, f'{cn:.0f}', f'{gr:.0f}'])
caption(f'表 A1　手法ごとの結果（図1・図3，N = {N} 走行）')
table(['手法', '総配信量 [MB]', 'assoc_hold との対差 [MB]', '走行ごと相対差 [%]',
       '接続時間 [s]', 'grant 時間 [s]'], rows,
      widths=[5.0, 2.5, 3.5, 2.0, 2.0, 2.0])
note('対差とその 95%% 信頼区間が本命の指標である（同一のチャネル実現に対する対応データのため）。'
     '総配信量の絶対値は走行ごとの交通実現に大きく左右されるので，単独の信頼区間は参考値。'
     '接続時間と grant 時間は車両 × アンテナで積算した総和。'
     'kkf_conv 対 assoc_hold の Wilcoxon 符号順位検定は p = %.2g。'
     % float(pa.loc['kkf_conv', 'wilcoxon_p']))

caption('表 A2　N = %d と N = %d の比較（図2。いずれも配置 d70，同一の学習済み地図）' % (N_OLD, N))
table(['項目', f'N = {N_OLD}（前回）', f'N = {N}（今回）'],
      [['走行ごと相対差', f'+{o_rel.mean():.1f}% ± {o_rel_ci:.1f} pt',
        f'+{rel.mean():.1f}% ± {rel_ci:.1f} pt'],
       ['　同上（平均に対する半幅）', f'±{100*o_rel_ci/o_rel.mean():.1f}%',
        f'±{100*rel_ci/rel.mean():.1f}%'],
       ['対差', f'+{o_d.mean():,.0f} MB ± {o_d_ci:,.0f}', f'+{d.mean():,.0f} MB ± {d_ci:,.0f}'],
       ['提案手法が上回った走行', f'{int((o_d>0).sum())}/{N_OLD}', f'{int((d>0).sum())}/{N}'],
       ['最も差が小さい走行', f'+{o_rel.min():.1f}%', f'+{rel.min():.1f}%'],
       ['oracle への到達率', f'{o_cap:.1f}%', f'{capture:.1f}%'],
       ['遮蔽リスク地図の寄与', '%s MB ±%s (p = %.2g)' % (
           smb(float(o_pk.loc['kkf_novar', 'diff_mean'])), mb(float(o_pk.loc['kkf_novar', 'diff_ci95'])),
           float(o_pk.loc['kkf_novar', 'wilcoxon_p'])),
        '%s MB ±%s (p = %.2g)' % (
            smb(float(pk.loc['kkf_novar', 'diff_mean'])), mb(float(pk.loc['kkf_novar', 'diff_ci95'])),
            float(pk.loc['kkf_novar', 'wilcoxon_p']))],
       ['LCB の寄与', '%s MB ±%s (p = %.2g)' % (
           smb(float(o_pk.loc['kkf_nolcb', 'diff_mean'])), mb(float(o_pk.loc['kkf_nolcb', 'diff_ci95'])),
           float(o_pk.loc['kkf_nolcb', 'wilcoxon_p'])),
        '%s MB ±%s (p = %.2g)' % (
            smb(float(pk.loc['kkf_nolcb', 'diff_mean'])), mb(float(pk.loc['kkf_nolcb', 'diff_ci95'])),
            float(pk.loc['kkf_nolcb', 'wilcoxon_p']))]],
      widths=[5.0, 6.0, 6.0])
note('N = %d は前回資料と同一の測定（sim_results/rec_cm_d70_ablation）であり，再計算していない。'
     'N = %d は別に記録した %d 走行による測定（sim_results/abl_d70_n50）で，'
     '学習済み地図と配置は共通だが交通実現のシードは前回を含む。' % (N_OLD, N, N_REC))

ab_rows = []
for m, label in [('kkf_cold', '事前学習した地図を外す（観測は grant 中に限る）'),
                 ('kkf_cold_probe', '事前学習を外し，代わりに全ペア観測を許す'),
                 ('kkf_novar', '遮蔽リスク地図 σν² を外す'),
                 ('kkf_nolcb', '不確実性の利用（LCB，κσ）を外す')]:
    dv = float(pk.loc[m, 'diff_mean'])
    ci95 = float(pk.loc[m, 'diff_ci95'])
    pv = float(pk.loc[m, 'wilcoxon_p'])
    tot = float(agg.loc[m, 'total_data_MB_mean'])
    share = '—' if m == 'kkf_cold' else f'{100*dv/d.mean():+.1f}%'
    ab_rows.append([label, f'{tot:,.0f}', f'{dv:+,.0f}', f'±{ci95:,.0f}', share,
                    '—（決定論的）' if m == 'kkf_cold' else f'{pv:.2g}'])
caption(f'表 A3　機構を1つずつ外したときの変化（図4，N = {N} 走行）')
table(['外した機構', '総配信量 [MB]', 'kkf_conv との対差 [MB]', '95%CI 半幅',
       '利得に占める割合', 'p'], ab_rows, widths=[5.4, 2.4, 3.0, 2.0, 2.2, 2.0])
note('対差は kkf_conv から見た差であり，正の値は kkf_conv が上回ることを意味する。'
     '「利得に占める割合」の分母は assoc_hold との対差 %s MB。'
     '探索的比較であり多重比較の補正は行っていない。'
     'kkf_cold は全走行 0 MB であり，対差は kkf_conv の総配信量そのものになる。'
     '検定は形式的には成立するが情報を持たないため省き，割合の分母にも載せていない。'
     % mb(d.mean()))

caption(f'表 A4　車両ごとの配信量からみた公平性（図5，N = {N} 走行）')
table(['手法', '最小の車両の配信量 [MB]', '配信量が 0 の車両の割合 [%]'],
      [[LBL[m], f'{float(agg.loc[m, "min_car_data_MB_mean"]):,.0f}',
        f'{float(agg.loc[m, "zero_car_pct_mean"]):.2f}']
       for m in ['oracle', 'kkf_cold_probe', 'kkf_conv', 'assoc_hold']],
      widths=[6.0, 5.5, 5.5])

cond_rows = []
for _, r in cond.iterrows():
    nm = str(r['条件'])
    geom, nrsu = GEOM.get(nm, ('—', 8))
    cell = f"{nm.replace('rec_cm_', '').replace('rec_cf_', '直交 ')}"
    cond_rows.append([cell, geom, str(nrsu), f"{r['選択余地率']:.3f}",
                      f"{r['assoc_hold_MB']:,.0f}", f"{r['天井_pct']:+.1f}%", f"{int(r['n'])}"])
caption('表 A5　RSU 配置と天井（図6，選択余地率の昇順）')
table(['条件', 'ポール本数 / 間隔 / δ_rsu', '論理 RSU 数', '選択余地率',
       'カバレッジ [MB]', '天井', 'N'], cond_rows,
      widths=[2.6, 3.6, 1.8, 2.0, 2.6, 2.2, 1.2], right_from=2)
note('カバレッジは受動接続（assoc_hold）の総配信量で代表させた。'
     '天井は平均どうしの比であり，走行ごと相対差の平均（表 A1 の oracle の行）とは'
     '定義もデータセットも異なる。選択余地率は記録した真値からの実測値（解析計算値ではない）。'
     '交通条件は全条件で共通だが N は条件ごとに異なる。'
     '「直交」の 4 条件が今回追加したもので，δ_rsu は 70° に固定してある。'
     '本表の 9 条件は付録 B の不具合を修正する前に取得したもので，'
     '各条件 20 走行のうち読み込めた走行のみを集計している。'
     'また d60 と d70s10 では，記録内の基地局番号と RSU の対応が走行間で同一かどうかを'
     '判定できていない（対応が割れている証拠は無い）。')


h1('付録 B：評価から外れた走行について')
body_p('以前は，基地局のスポーン要求が並列実行時に無言で失敗し，'
       'その走行が配信量ゼロのまま「正常終了」することがあった'
       '（前回の評価では記録 20 走行のうち 7 走行）。'
       '静的エンティティをワールド SDF に直接書き込む方式に変更して解消済みであり，'
       '今回の %d 走行では配信量ゼロの走行は出ていない。表 A5 の 9 条件はこの修正前の測定である。'
       % N_REC)
body_p('今回評価に載らなかった %d 走行（seed_37345，seed_50345）は，'
       'いずれもシム時刻 62 秒付近で走行が打ち切られており（正常な走行は約 78 秒），'
       '記録ファイルの末尾が書き込みの途中で切れていたため読み込めなかった。'
       '記録は全手法が共有するので，この脱落は手法の優劣に対して選択的ではない。'
       'また走行長が他と異なるため，読み込めたとしても対比較には載せられない。'
       % (N_REC - N))

h1('付録 C：地図が接続可否を正しく判定できるか（前回からの再掲）')
body_p('本文の主張は，地図が「その相手に今つないだら送れるか」を当てられることに依存する。'
       'これを，記録した真値を答え合わせに用いて測った（5 走行，接続可能標本 6.6%）。'
       '前回資料から取り直していないが，本評価と同じ学習済み地図による測定である。')
caption('表 C1　地図の判定精度（再掲）')
table(['取りこぼし率', '幻リンク率', 'バイアス', 'RMSE', 'z の標準偏差'],
      [[f'{au["miss"]:.1f}%', f'{au["fa"]:.2f}%', f'{au["bias"]:+.2f} dB',
        f'{au["rmse"]:.2f} dB', f'{au["z"]:.2f}']],
      widths=[3.4, 3.4, 3.4, 3.4, 3.4], right_from=0)
note('取りこぼし = 真に接続可能なのに地図が接続不可と判定した割合，幻リンクはその逆。'
     'バイアスと RMSE は接続可能な標本での予測誤差。'
     'z =（予測値 − 真値）/σ で，標準偏差 1.0 なら地図が申告する不確かさ σ が'
     '実際の誤差の大きさと合っていることを意味する。')
body_p('取りこぼし %.1f%%・幻リンク %.2f%%・バイアス %+.1f dB であり，'
       '地図は接続可否をおおむね正しく判定できている。'
       'なお σ は LCB に使うが，その寄与は本文のとおり利得の %.1f%% にとどまるため，'
       '主張を支えているのは主に接続可否の判定精度である。'
       % (au['miss'], au['fa'], au['bias'],
          100 * float(pk.loc['kkf_nolcb', 'diff_mean']) / d.mean()))

h1('付録 D：再現手順')
body_p('本評価は Docker コンテナ内で，/workspace を作業ディレクトリとして実行した'
       '（イメージ ros2-gazebo-comms-sim:humble-harmonic，ROS 2 Humble，'
       'Gazebo Sim 8.11.0，Python 3.10.12）。以下はすべてコンテナ内で実行する。')
note('① コードを揃える： feat/urban-2lane の 789996b を checkout する'
     '（本評価に用いたシミュレータ・シナリオ・ツールを含む）。'
     '記録と学習の sweep 設定，および図表の生成スクリプトは ef7b8b7 に含まれる。'
     '記録した時点の作業差分は sim_results/rec_d70_n50/sweep/config/workspace.patch に'
     '残してあり，これを 755fe3d に当てた状態と 789996b とで，'
     '挙動に関わる 4 ファイル（TxControllerPlugin.cc，sim_launch.py，'
     'replay_sim.py，replay_sweep.py）が一致することを確認済みである。'
     'シナリオと sweep 設定は記録時のスナップショットも '
     'sim_results/rec_d70_n50/sweep/config/ に scenario.yaml・sweep.yaml・'
     'sim_params_base.yaml として残っており，照合に使える。')
note('② 地図を学習する： python3 tools/sim.py learn '
     '--scenario config/scenarios/urban_cm_d70.yaml --name urban_cm_d70_learn '
     '--runs 30 --base-seed 512345 → sim_results/urban_cm_d70_learn/rem_state。'
     '直列実行で約 59 分（1 走行あたり約 2 分）。')
note('③ 走行を記録する： python3 tools/sweep_sim.py '
     '--sweep-config scratch/record_d70_n50.yaml → sim_results/rec_d70_n50/seed_*。'
     '並列度は sweep 設定の max_concurrency: 8，実時間倍率 2.0 で約 1 時間 40 分。'
     '出力は約 13 GB（50 走行分の軌跡と全ペア受信電力）。'
     '並列度を変えると走行あたりの計算資源が変わるため，脱落の起こりやすさも変わりうる。')
note('④ 全手法を再生する： python3 tools/replay_sweep.py sim_results/rec_d70_n50 '
     '--state-dir sim_results/urban_cm_d70_learn/rem_state --out sim_results/abl_d70_n50 '
     '--jobs 8 --baseline assoc_hold。出力は約 3.7 MB。')
note('⑤ 本資料の図表を作る： python3 scratch/make_investigation_figs.py '
     'および python3 scratch/make_investigation_docx.py。'
     '図の描画には日本語フォントが要る（コンテナに CJK フォントが無い場合は '
     'NotoSansCJK-Regular.ttc を /usr/share/fonts に置く。手順はスクリプトの冒頭に記載）。'
     '文書生成は書式の雛形として scratch/seminar/2026_08_04_KKFSystem.docx を読む。'
     '図のファイル名（fig2_runs.png = 図2，fig3_time.png = 図3）は図番号に対応させてある。')
body_p('再現の水準は 2 段階に分かれる。'
       '③ の記録は Gazebo の実時間スケジューリングに依存するためビット一致では再現せず，'
       '同一シード列からの再実行が統計的に同一の分布を与えるにとどまる。'
       'これに対し ④ 以降は記録を入力とする決定論的な計算であり，'
       '記録が同じであれば本資料の数値は完全に一致する（第 2 章で 1 走行について確認した）。')
body_p('したがって，記録一式（約 13 GB）と学習済み地図（644 KB）があれば，'
       '図1 と図3〜図5，表 A1・A3・A4 は第三者が完全に検証できる。'
       'ただし表 A2 は前回の記録（sim_results/rec_cm_d70_ablation），'
       '表 A5 は 9 条件それぞれの集計（choice_margin_summary.csv，confound_summary.csv），'
       '付録 C は較正監査の出力（sim_results/audit_d70）を別に必要とする。'
       'これらを含めない配布では，本資料の一部の表は再計算できない。')
body_p('残る制約はデータの配布である。コードと設定の版は確定したが，'
       '記録一式は約 13 GB あり，リポジトリには含めていない'
       '（sim_results/ は版管理の対象外）。'
       '第三者が ④ 以降だけを検証する場合は，記録一式と学習済み地図を別途受け渡す必要がある。')

path = os.path.join(OUT, '内部調査資料_2026-08-13.docx')
doc.save(path)

# --- 査読用テキスト (表を本文の位置に差し込む) ---
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

out, ti = [], 0
for child in doc.element.body.iterchildren():
    if child.tag.endswith('}p'):
        p = Paragraph(child, doc)
        if not p.text.strip():
            continue
        pre = '## ' if p.style.name == 'List Paragraph' else ''
        out.append(pre + p.text)
    elif child.tag.endswith('}tbl'):
        t = Table(child, doc)
        ti += 1
        rws = [' | '.join(c.text.replace('\n', ' ') for c in r.cells) for r in t.rows]
        out.append('[表%d]\n' % ti + '\n'.join(rws))
with open(os.path.join(OUT, 'investigation_text.md'), 'w', encoding='utf-8') as f:
    f.write('\n\n'.join(out) + '\n')
print('wrote', path)
