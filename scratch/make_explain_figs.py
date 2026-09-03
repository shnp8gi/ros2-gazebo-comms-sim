#!/usr/bin/env python3
"""内部調査資料の説明図 (A/B/C) を PowerPoint で作る。

図形はすべてネイティブなので PowerPoint 上で編集できる。
スタイルは既存の scratch/review/system_flowchart.pptx に合わせる
(白黒・MS Mincho・線幅 1pt・スライド 33.87 x 19.05 cm)。
"""
import os
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   'review', 'figures_explain.pptx')
LATIN = 'Times New Roman'   # 英数字
EA = 'MS Mincho'            # 日本語


def set_font(run, size, bold, color):
    """英数字は Times New Roman、日本語は MS 明朝に分ける。

    python-pptx の font.name は latin しか設定しないため、ea/cs は XML を直接書く。
    """
    from pptx.oxml.ns import qn as _qn
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag, face in (('a:latin', LATIN), ('a:ea', EA), ('a:cs', EA)):
        for old in rPr.findall(_qn(tag)):
            rPr.remove(old)
        el = rPr.makeelement(_qn(tag), {'typeface': face})
        rPr.append(el)
K = RGBColor(0, 0, 0)
W_ = RGBColor(0xFF, 0xFF, 0xFF)
G = RGBColor(0x80, 0x80, 0x80)

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.87), Cm(19.05)
BLANK = prs.slide_layouts[6]



M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
A14_NS = 'http://schemas.microsoft.com/office/drawing/2010/main'


def omath(par, text):
    """段落に PowerPoint の数式オブジェクト (挿入>数式) を差し込む。

    python-pptx に数式の API が無いため、a14:m でくるんだ m:oMath を直接書く。
    これは PowerPoint が「数式」として認識・編集できる形式。
    """
    import re as _re
    from lxml import etree
    m = etree.SubElement(par._p, '{%s}m' % A14_NS)
    om = etree.SubElement(m, '{%s}oMath' % M_NS)
    # 変数(1文字の英字/ギリシャ文字)は斜体、それ以外は立体で組む
    for tok in _re.findall(r'[A-Za-zα-ωΑ-Ω]|[^A-Za-zα-ωΑ-Ω]+', text):
        r = etree.SubElement(om, '{%s}r' % M_NS)
        rPr = etree.SubElement(r, '{%s}rPr' % M_NS)
        if not _re.fullmatch(r'[A-Za-zα-ωΑ-Ω]', tok):
            etree.SubElement(rPr, '{%s}nor' % M_NS)     # 立体 (normal text)
        apr = etree.SubElement(r, '{%s}rPr' % A_NS)
        for tag, face in (('a:latin', LATIN), ('a:ea', EA)):
            apr.append(apr.makeelement(qn(tag), {'typeface': face}))
        t = etree.SubElement(r, '{%s}t' % M_NS)
        t.text = tok
    return m


def mathbox(sl, x, y, w, h, text, size=11, align=PP_ALIGN.CENTER):
    """数式だけを持つテキストボックス。"""
    s = sl.shapes.add_textbox(x, y, w, h)
    f = s.text_frame
    f.word_wrap = False
    f.margin_left = f.margin_right = Cm(0.03)
    f.margin_top = f.margin_bottom = 0
    f.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = f.paragraphs[0]
    p.alignment = align
    omath(p, text)
    return s

def tb(sl, x, y, w, h, text, size=11, bold=False, color=K,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    s = sl.shapes.add_textbox(x, y, w, h)
    f = s.text_frame
    f.word_wrap = True
    f.margin_left = f.margin_right = Cm(0.05)
    f.margin_top = f.margin_bottom = 0
    f.vertical_anchor = anchor
    for i, ln in enumerate(text.split('\n')):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        set_font(r, size, bold, color)
    return s


def box(sl, x, y, w, h, text='', size=11, bold=False, fill=W_,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=K, lw=1.0, dash=None):
    s = sl.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(lw)
    if dash: s.line.dash_style = dash
    s.shadow.inherit = False
    f = s.text_frame; f.word_wrap = True
    f.margin_left = f.margin_right = Cm(0.08)
    f.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text.split('\n')):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ln
        set_font(r, size, bold, K)
    return s


def line(sl, x1, y1, x2, y2, dash=None, lw=1.0, arrow=False, color=K):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if dash: c.line.dash_style = dash
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'),
                                 {'type': 'triangle', 'w': 'med', 'h': 'med'}))
    return c


def title(sl, t):
    tb(sl, Cm(0.8), Cm(0.5), Cm(32), Cm(1.0), t, size=16, bold=True,
       align=PP_ALIGN.LEFT)


def note(sl, t):
    tb(sl, Cm(0.8), Cm(17.8), Cm(32), Cm(0.9), t, size=9.5, color=G,
       align=PP_ALIGN.LEFT)


# ==================================================== 図A 問題設定
sl = prs.slides.add_slide(BLANK)
title(sl, '図A　問題設定：排他接続と、観測できないペア')

RSU_Y, VEH_Y = Cm(3.4), Cm(11.0)
RW, RH = Cm(3.6), Cm(1.4)
VW, VH = Cm(3.6), Cm(1.4)
rsu_x = [Cm(4.5), Cm(14.5), Cm(24.5)]
veh_x = [Cm(3.0), Cm(11.5), Cm(20.0), Cm(27.0)]

# 道路
box(sl, Cm(1.6), Cm(10.4), Cm(30.6), Cm(2.6), '', shape=MSO_SHAPE.RECTANGLE,
    line=G)
tb(sl, Cm(1.6), Cm(13.1), Cm(30.6), Cm(0.6), '走行方向 →', size=10, color=G,
   align=PP_ALIGN.RIGHT)

for i, x in enumerate(rsu_x):
    box(sl, x, RSU_Y, RW, RH, f'RSU {i+1}', size=12, bold=True)
for i, x in enumerate(veh_x):
    box(sl, x, VEH_Y, VW, VH, f'車両 {chr(65+i)}', size=12)

def mid(x, w): return x + w // 2

# 接続中 (実線・太)
line(sl, mid(rsu_x[0], RW), RSU_Y + RH, mid(veh_x[0], VW), VEH_Y, lw=2.5)
line(sl, mid(rsu_x[1], RW), RSU_Y + RH, mid(veh_x[2], VW), VEH_Y, lw=2.5)
# 同時接続できない (破線)
line(sl, mid(rsu_x[0], RW), RSU_Y + RH, mid(veh_x[1], VW), VEH_Y,
     dash=MSO_LINE_DASH_STYLE.DASH, color=G)
# 未観測 (破線)
for rx, vx in [(rsu_x[1], veh_x[1]), (rsu_x[2], veh_x[2]), (rsu_x[2], veh_x[3]),
               (rsu_x[1], veh_x[3])]:
    line(sl, mid(rx, RW), RSU_Y + RH, mid(vx, VW), VEH_Y,
         dash=MSO_LINE_DASH_STYLE.DASH, color=G)

# 凡例
lx, ly = Cm(1.6), Cm(14.4)
line(sl, lx, ly + Cm(0.3), lx + Cm(1.6), ly + Cm(0.3), lw=2.5)
tb(sl, lx + Cm(1.8), ly, Cm(13), Cm(0.6),
   '接続中のペア：受信電力を継続して測れる（1 RSU につき 1 台のみ）',
   size=11, align=PP_ALIGN.LEFT)
line(sl, lx, ly + Cm(1.2), lx + Cm(1.6), ly + Cm(1.2),
     dash=MSO_LINE_DASH_STYLE.DASH, color=G)
tb(sl, lx + Cm(1.8), ly + Cm(0.9), Cm(16), Cm(0.6),
   '未接続のペア：割当の判断に足る受信電力を測れない',
   size=11, align=PP_ALIGN.LEFT)

box(sl, Cm(19.5), Cm(14.3), Cm(12.7), Cm(2.6),
    'どの車両にどの RSU をいつ与えるかを決めるには、\n'
    '測っていない相手の電波状況を何らかの方法で補う必要がある',
    size=11.5, shape=MSO_SHAPE.RECTANGLE)
note(sl, '802.15.3e のペアネット方式。規格の受動アソシエーションで分かるのは'
         '「つなげる見込みがあるか」の粗い判定にとどまる。')

# ==================================================== 図B 配置 (上面図+断面図)
sl = prs.slides.add_slide(BLANK)
title(sl, '図B　シミュレーションの配置')

# --- 上面図
tb(sl, Cm(0.8), Cm(1.55), Cm(4.2), Cm(0.55), '(a) 上面図', size=12, bold=True,
   align=PP_ALIGN.LEFT)
X0, SCALE = Cm(3.0), Cm(0.33)          # x=-40m を X0 に、1 m = 0.33 cm
def px(m): return X0 + Emu(int(SCALE * (m + 40)))
TOP_Y = Cm(2.65)
LANE_H = Cm(1.15)
# 車線
box(sl, px(-40), TOP_Y + Cm(1.6), px(40) - px(-40), LANE_H, '',
    shape=MSO_SHAPE.RECTANGLE, line=G)
box(sl, px(-40), TOP_Y + Cm(2.75), px(40) - px(-40), LANE_H, '',
    shape=MSO_SHAPE.RECTANGLE, line=G)
mathbox(sl, Cm(0.4), TOP_Y + Cm(1.6), Cm(2.5), LANE_H, 'y = +1.75 m',
        size=9, align=PP_ALIGN.RIGHT)
mathbox(sl, Cm(0.4), TOP_Y + Cm(2.75), Cm(2.5), LANE_H, 'y = −1.75 m',
        size=9, align=PP_ALIGN.RIGHT)
# 歩道側の RSU ポール (y=6.0)
for xm in (-30, -10, 10, 30):
    box(sl, px(xm) - Cm(0.35), TOP_Y, Cm(0.7), Cm(0.7), '',
        shape=MSO_SHAPE.OVAL)
    mathbox(sl, px(xm) - Cm(1.4), TOP_Y - Cm(0.62), Cm(2.8), Cm(0.5),
            f'x = {xm} m', size=8.5)
    # 上流/下流の2面
    line(sl, px(xm), TOP_Y + Cm(0.7), px(xm) - Cm(1.2), TOP_Y + Cm(1.6),
         arrow=True)
    line(sl, px(xm), TOP_Y + Cm(0.7), px(xm) + Cm(1.2), TOP_Y + Cm(1.6),
         arrow=True)
mathbox(sl, Cm(0.4), TOP_Y, Cm(2.5), Cm(0.7), 'y = 6.0 m', size=9,
        align=PP_ALIGN.RIGHT)
# 路上駐車 (y=4.0)
for xm, lab in ((-9, '路上駐車 3.2 m'), (7, '路上駐車 2.4 m')):
    box(sl, px(xm) - Cm(1.0), TOP_Y + Cm(1.02), Cm(2.0), Cm(0.45), '',
        shape=MSO_SHAPE.RECTANGLE, fill=RGBColor(0xD9, 0xD9, 0xD9))
    # 指向を示す矢印と重ならないよう、注記は車線帯の下へ出す
    tb(sl, px(xm) - Cm(2.0), TOP_Y + Cm(4.0), Cm(4.0), Cm(0.5), lab, size=8.5)
    line(sl, px(xm), TOP_Y + Cm(1.47), px(xm), TOP_Y + Cm(4.0), color=G,
         dash=MSO_LINE_DASH_STYLE.DASH)
# 車両
for xm in (-34, -22, -14, -2, 6, 18, 26, 34):
    box(sl, px(xm) - Cm(0.7), TOP_Y + Cm(1.75), Cm(1.4), Cm(0.85), '',
        shape=MSO_SHAPE.RECTANGLE)
for xm in (-30, -18, -6, 10, 22, 32):
    box(sl, px(xm) - Cm(0.7), TOP_Y + Cm(2.9), Cm(1.4), Cm(0.85), '',
        shape=MSO_SHAPE.RECTANGLE)
tb(sl, Cm(24.0), TOP_Y + Cm(4.1), Cm(9), Cm(0.6),
   'ポール 4 本 × 上流/下流 2 面 = 論理 RSU 8 基（間隔 20 m）', size=10,
   align=PP_ALIGN.RIGHT)

# --- 断面図
tb(sl, Cm(0.8), Cm(8.6), Cm(10), Cm(0.6), '(b) 断面図（高さ関係）', size=12,
   bold=True, align=PP_ALIGN.LEFT)
GY = Cm(16.0)                            # 地面
HS = Cm(1.45)                            # 1 m の高さ
def py(m): return GY - Emu(int(HS * m))
def cx(ym): return Cm(6.2) + Emu(int(Cm(1.62) * (6.0 - ym)))   # y=6 を左に

line(sl, Cm(3.05), GY, Cm(21.0), GY, lw=1.5)
tb(sl, Cm(3.05), GY + Cm(0.15), Cm(2.4), Cm(0.5), '地面', size=9, color=G,
   align=PP_ALIGN.LEFT)
# 高さ目盛 (左端に寄せる)
for hm in (1, 2, 3):
    line(sl, Cm(3.05), py(hm), Cm(3.3), py(hm), color=G)
    mathbox(sl, Cm(1.55), py(hm) - Cm(0.28), Cm(1.35), Cm(0.56), f'{hm} m',
            size=8.5, align=PP_ALIGN.RIGHT)

# RSU ポール (y = 6.0 m)
RSUX = cx(6.0)
line(sl, RSUX, GY, RSUX, py(2.5), lw=2.0)
box(sl, RSUX - Cm(0.7), py(2.5) - Cm(0.3), Cm(1.4), Cm(0.6), '',
    shape=MSO_SHAPE.RECTANGLE)
tb(sl, RSUX - Cm(1.6), py(2.5) - Cm(1.35), Cm(3.5), Cm(0.55),
   'RSU アンテナ', size=10)
mathbox(sl, RSUX - Cm(1.6), py(2.5) - Cm(0.85), Cm(3.5), Cm(0.5), '2.5 m',
        size=10)

# 手前車線の車 (y = +1.75 m)
NX = cx(1.75)
box(sl, NX - Cm(1.35), py(1.5), Cm(2.7), Emu(int(HS * 1.5)), '',
    shape=MSO_SHAPE.RECTANGLE)
tb(sl, NX - Cm(1.35), py(1.5) + Cm(0.15), Cm(2.7), Cm(0.5), '乗用車', size=9)
mathbox(sl, NX - Cm(1.35), py(1.5) + Cm(0.62), Cm(2.7), Cm(0.45), '1.5 m',
        size=9)

# 奥車線の車 (y = −1.75 m) と車載アンテナ
FX = cx(-1.75)
box(sl, FX - Cm(1.35), py(1.5), Cm(2.7), Emu(int(HS * 1.5)), '',
    shape=MSO_SHAPE.RECTANGLE)
box(sl, FX - Cm(0.22), py(1.35) - Cm(0.22), Cm(0.44), Cm(0.44), '',
    shape=MSO_SHAPE.OVAL, fill=K)
tb(sl, FX - Cm(0.6), py(1.35) - Cm(1.35), Cm(5.6), Cm(0.55),
   '車載アンテナ', size=10, align=PP_ALIGN.LEFT)
mathbox(sl, FX - Cm(0.6), py(1.35) - Cm(0.85), Cm(5.6), Cm(0.5), '1.35 m',
        size=10, align=PP_ALIGN.LEFT)

# 見通し線
line(sl, FX, py(1.35), RSUX, py(2.5), lw=2.0, dash=MSO_LINE_DASH_STYLE.DASH)
# 手前車線を横切る高さ
line(sl, NX, py(1.87), NX, py(0), color=G, dash=MSO_LINE_DASH_STYLE.DASH)
box(sl, NX - Cm(0.25), py(1.87) - Cm(0.25), Cm(0.5), Cm(0.5), '',
    shape=MSO_SHAPE.OVAL, fill=K)
tb(sl, NX - Cm(2.9), py(1.87) - Cm(2.35), Cm(5.8), Cm(0.55),
   '見通し線が手前車線を', size=10.5, bold=True)
tb(sl, NX - Cm(2.9), py(1.87) - Cm(1.85), Cm(5.8), Cm(0.55),
   '横切る高さ', size=10.5, bold=True)
mathbox(sl, NX - Cm(2.9), py(1.87) - Cm(1.32), Cm(5.8), Cm(0.5), '1.87 m',
        size=11.5)

# --- 車高の比較 (右側に独立した領域として置く)
BX, BY = Cm(23.2), GY
tb(sl, BX, Cm(8.6), Cm(9.6), Cm(0.6), '遮蔽するかどうかは車高で決まる',
   size=11, bold=True, align=PP_ALIGN.LEFT)
for i, (hm, jp, num) in enumerate([(1.5, '乗用車', '1.5 m'),
                                   (1.9, 'ミニバン', '1.9 m'),
                                   (3.2, 'バス', '3.2 m'),
                                   (3.8, 'トラック', '3.8 m')]):
    x = BX + Cm(2.35) * i
    box(sl, x, BY - Emu(int(HS * hm)), Cm(1.5), Emu(int(HS * hm)), '',
        shape=MSO_SHAPE.RECTANGLE,
        fill=(W_ if hm < 1.87 else RGBColor(0xD9, 0xD9, 0xD9)))
    tb(sl, x - Cm(0.4), BY + Cm(0.15), Cm(2.3), Cm(0.5), jp, size=9)
    mathbox(sl, x - Cm(0.4), BY + Cm(0.62), Cm(2.3), Cm(0.45), num, size=9)
line(sl, BX - Cm(0.5), BY - Emu(int(HS * 1.87)), BX + Cm(9.0),
     BY - Emu(int(HS * 1.87)), dash=MSO_LINE_DASH_STYLE.DASH, lw=1.5)
mathbox(sl, BX + Cm(6.3), BY - Emu(int(HS * 1.87)) - Cm(0.55), Cm(2.7),
        Cm(0.5), '1.87 m', size=9.5, align=PP_ALIGN.RIGHT)
tb(sl, BX - Cm(0.5), BY - Emu(int(HS * 1.87)) - Cm(1.05), Cm(9.5), Cm(0.5),
   'この線を超える車高が遮蔽する', size=9.5, align=PP_ALIGN.LEFT)
line(sl, GY - GY + Cm(21.6), Cm(9.4), Cm(21.6), Cm(16.6), color=G)

note(sl, '断面図は奥車線の車両から RSU への見通しを示す。手前車線を横切る高さは '
         '1.87 m で、乗用車（1.5 m）では届かない。本評価では走行車両を遮蔽体として'
         '扱わず、電波をさえぎるのは位置の固定された路上駐車 2 台のみである。')

# ==================================================== 図C 評価パイプライン
sl = prs.slides.add_slide(BLANK)
title(sl, '図C　評価の手続き：1 回記録し、手法だけを差し替えて事後再計算する')

Y1, BH = Cm(3.2), Cm(2.0)
box(sl, Cm(1.2), Y1, Cm(7.6), BH,
    'Gazebo で走行を実行\n（50 走行・手法に依存しない）', size=11.5, bold=True)
line(sl, Cm(8.8), Y1 + BH // 2, Cm(10.4), Y1 + BH // 2, arrow=True, lw=1.5)
box(sl, Cm(10.4), Y1, Cm(8.4), BH,
    '記録\n・軌跡 poses.csv\n・全ペアの受信電力 pairs/*.csv', size=11)
line(sl, Cm(18.8), Y1 + BH // 2, Cm(20.4), Y1 + BH // 2, arrow=True, lw=1.5)
box(sl, Cm(20.4), Y1, Cm(7.8), BH,
    '欠落検査\nverify_recording.py', size=11)

# 評価窓
box(sl, Cm(10.4), Cm(6.2), Cm(8.4), Cm(1.5),
    '評価窓を揃える\n走り出しから 35 秒', size=11)
line(sl, Cm(14.6), Y1 + BH, Cm(14.6), Cm(6.2), arrow=True)

# 事後再計算
Y2 = Cm(9.0)
arms = ['受動接続\nassoc_hold', '提案手法\nrem_conv', '地図なし\nrem_cold',
        '機構を1つ外す\nrem_nolcb / novar /\nnokrig / frozen',
        '理想反応型\noracle']
AW = Cm(5.6)
for i, a in enumerate(arms):
    x = Cm(1.2) + (AW + Cm(0.65)) * i
    box(sl, x, Y2, AW, Cm(2.4), a, size=10.5)
    line(sl, x + AW // 2, Cm(7.7) if i == 1 else Cm(7.7), x + AW // 2, Y2,
         arrow=True, dash=MSO_LINE_DASH_STYLE.DASH, color=G)
line(sl, Cm(14.6), Cm(7.7), Cm(4.0), Cm(7.7), color=G)
line(sl, Cm(14.6), Cm(7.7), Cm(29.0), Cm(7.7), color=G)
tb(sl, Cm(1.2), Cm(11.6), Cm(31.4), Cm(0.7),
   '同一の記録＝同一の交通・同一のチャネル実現の上で、通信とスケジューリングだけを再計算する',
   size=11.5, bold=True)

line(sl, Cm(16.9), Cm(12.4), Cm(16.9), Cm(13.4), arrow=True, lw=1.5)
box(sl, Cm(9.4), Cm(13.4), Cm(15.0), Cm(1.8),
    '走行ごとに対にして比較（対差・Wilcoxon 符号順位検定）', size=12, bold=True)

box(sl, Cm(1.2), Cm(15.6), Cm(14.6), Cm(1.7),
    'Gazebo は不要。手法を増やしても再記録は要らない', size=11,
    shape=MSO_SHAPE.RECTANGLE)
box(sl, Cm(17.6), Cm(15.6), Cm(14.6), Cm(1.7),
    '再計算は決定論的で、走行ごとの交通の当たり外れが打ち消される', size=11,
    shape=MSO_SHAPE.RECTANGLE)
note(sl, '通信計算は車両の運動に影響しないため、軌跡と全ペアの受信電力を一度記録すれば、'
         '手法・設定を変えた評価は単一プロセスで再現できる。')

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print(f'wrote {OUT}')
